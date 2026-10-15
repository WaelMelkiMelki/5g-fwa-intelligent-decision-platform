# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu

PATH = r"C:\Users\w50057230\Downloads\presnetation slides\FWA_Jury_Deck_FULL.pptx"
prs = Presentation(PATH)
EMU_IN = 914400
W = prs.slide_width / EMU_IN
H = prs.slide_height / EMU_IN
print("canvas:", round(W, 2), "x", round(H, 2), "| slides:", len(prs.slides._sldIdLst))

def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh

issues = 0
for idx, slide in enumerate(prs.slides, 1):
    pics = 0
    for sh in walk(slide.shapes):
        if sh.shape_type == 13:  # picture
            pics += 1
        # bounds check
        try:
            l, t = sh.left / EMU_IN, sh.top / EMU_IN
            w, h = sh.width / EMU_IN, sh.height / EMU_IN
        except TypeError:
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            if l < -0.05 or t < -0.05 or l + w > W + 0.05 or t + h > H + 0.05:
                print(f"  [S{idx}] OUT-OF-BOUNDS text: '{sh.text_frame.text[:40]}' at ({l:.2f},{t:.2f},{w:.2f},{h:.2f})")
                issues += 1
            # crude overflow estimate: chars per line at font size
            for p in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt: continue
                size = max((r.font.size.pt if r.font.size else 12) for r in p.runs)
                # est chars/line: width_in * 96 / (size * 0.55)
                cpl = max(int(w * 96 / (size * 0.55)), 1)
                lines = max(1, -(-len(txt) // cpl))
                need_h = lines * size * 1.28 / 72
                if need_h > h + 0.12 and h > 0.2:
                    print(f"  [S{idx}] OVERFLOW-RISK: '{txt[:45]}' size={size} box_h={h:.2f} need≈{need_h:.2f}")
                    issues += 1
    print(f"S{idx:02d}: shapes={len(list(walk(slide.shapes)))}, pics={pics}")

print("issues:", issues)
